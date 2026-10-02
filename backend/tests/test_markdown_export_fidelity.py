"""Regressions for Markdown fidelity in the Word and PDF exports.

Each test pins a defect seen in real exported answers: raw ``**[link](url)**`` text,
literal ``<br>`` in table cells, flattened nested lists, raw ``$\\rightarrow$``, misleading
``[x]`` check marks, a stranded first page before a diagram, a truncated footer model, and
Word lists that never restart numbering.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from PIL import Image
from pypdf import PdfReader
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate

from app import markdown_render as mr

STEPS = """Step 2:
1. Go to **Entra admin center**. Find the two policies:
   - **Use app-enforced restrictions**: the limited experience.
   - **Block access**: blocks desktop
     and mobile apps.
2. **Scope both** first.

3. Third, after a blank line
   ```powershell
   Get-SPOTenant
   ```
"""


def _png(width: int, height: int) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), "#6366f1").save(buf, format="PNG")
    return buf.getvalue()


# --------------------------------------------------------------------------- parsing


def test_nested_list_items_keep_their_children():
    blocks = mr.parse_blocks(STEPS)
    assert blocks[0] == ("p", "Step 2:")
    kind, items, start, children = blocks[1]
    assert (kind, start, len(blocks)) == ("ol", 1, 2)
    assert items[0].startswith("Go to **Entra admin center**")
    assert children[0] == [(
        "ul",
        ["**Use app-enforced restrictions**: the limited experience.",
         "**Block access**: blocks desktop and mobile apps."],
        [[], []],
    )]
    # A blank line between items does not end the list; indented code belongs to its item.
    assert items[2] == "Third, after a blank line"
    assert children[2] == [("code", "Get-SPOTenant", "powershell")]


def test_list_followed_by_other_list_type_is_a_new_list():
    blocks = mr.parse_blocks("1. one\n2. two\n- bullet")
    assert [b[0] for b in blocks] == ["ol", "ul"]


def test_emphasis_is_parsed_recursively():
    spans = mr._inline_spans("Open the **[Admin Center](https://a.example/x)** and **run `Get-X` now**")
    assert ("Admin Center", {"bold", "link"}, "https://a.example/x") in spans
    assert ("Get-X", {"bold", "code"}, None) in spans
    assert not any("**" in s[0] or "](" in s[0] or "`" in s[0] for s in spans)


def test_br_and_simple_tex_are_converted_but_dollar_amounts_are_not():
    spans = mr._inline_spans("**A**<br/>*(b)* costs $5 and $10; File $\\rightarrow$ Save; $x^2 \\le 4$")
    assert ("\n", {"br"}, None) in spans
    text = "".join(s[0] for s in spans)
    assert "costs $5 and $10" in text
    assert "File \u2192 Save" in text
    assert "x\u00b2 \u2264 4" in text
    # Unknown TeX stays as written rather than being mangled.
    assert "".join(s[0] for s in mr._inline_spans("$\\mathcal{O}(n)$")) == "$\\mathcal{O}(n)$"


# --------------------------------------------------------------------------- PDF


def test_symbol_fallbacks_are_unambiguous(monkeypatch):
    monkeypatch.setattr(mr, "_GLYPH_OK", lambda ch: ch.isascii())
    assert mr.pdf_safe("\u2705 yes \u274c no \u26a0\ufe0f careful") == "[OK] yes [X] no [!] careful"


def test_symbols_use_the_symbol_font_when_installed():
    if "symbol" not in mr.pdf_fonts():
        pytest.skip("no symbol font on this host")
    markup = mr.pdf_markup("\u2705 Option A \u274c \u26a0\ufe0f")
    assert markup.count('<font face="MCSymbol">') == 3
    assert "[OK]" not in markup and "\ufe0f" not in markup


def test_footer_keeps_the_model_and_marks_the_cut():
    fitted = mr._fit_footer_text(
        "Block Downloads on Unmanaged Devices \u00b7 gpt-5.6-sol-fast", "Helvetica", 7.5, 150
    )
    assert fitted.endswith("\u2026 \u00b7 gpt-5.6-sol-fast")
    assert fitted.startswith("Block Downloads")


def _pdf(md: str, diagrams=None) -> PdfReader:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER)
    doc.build(mr.markdown_pdf_flowables(
        md, getSampleStyleSheet()["BodyText"], diagrams=diagrams, content_width=doc.width,
    ))
    return PdfReader(io.BytesIO(buf.getvalue()))


def _diagram_after_heading(natural_width: float) -> PdfReader:
    code = "graph TD; A-->B"
    filler = "\n\n".join(["Filler paragraph text for the first part of the page. " * 5] * 3)
    md = f"{filler}\n\n### Architecture Flow\n\n```mermaid\n{code}\n```\n"
    return _pdf(md, [{"code": code, "data": _png(400, 500), "width": natural_width,
                      "height": natural_width * 1.25}])


def test_heading_before_a_diagram_does_not_strand_the_first_page():
    # Large labels: the diagram can shrink into the room left beside its heading.
    reader = _diagram_after_heading(400)
    assert len(reader.pages) == 1
    assert reader.pages[0].images
    assert "Architecture Flow" in reader.pages[0].extract_text()


def test_diagram_that_would_become_illegible_moves_with_its_heading():
    # A big layout squeezed into that room would print its labels below 5pt, so the
    # heading and diagram start the next page together instead.
    reader = _diagram_after_heading(1400)
    assert len(reader.pages) == 2
    assert not reader.pages[0].images and reader.pages[1].images
    assert "Architecture Flow" not in reader.pages[0].extract_text()
    assert "Architecture Flow" in reader.pages[1].extract_text()


def test_pdf_renders_nested_inline_markup_tables_and_lists():
    md = (
        "1. Sign in to the **[Admin Center](https://admin.example/sp)**.\n"
        "   - nested **`Set-SPOTenant`** step\n\n"
        "| Method | Licence |\n|---|---|\n| **A**<br>*(native)* | P1 $\\rightarrow$ E3 |\n"
    )
    reader = _pdf(md)
    text = " ".join(page.extract_text() for page in reader.pages)
    for raw in ("**", "](", "<br", "\\rightarrow", "`"):
        assert raw not in text
    assert "Admin Center" in text and "Set-SPOTenant" in text
    annots = [a.get_object() for p in reader.pages for a in p.get("/Annots", [])]
    assert any(a["/A"]["/URI"] == "https://admin.example/sp" for a in annots if "/A" in a)


# --------------------------------------------------------------------------- Word


def _docx(md: str, diagrams=None):
    doc = Document()
    mr.docx_page_setup(doc, "Title \u00b7 model")
    mr.render_markdown_docx(doc, md, diagrams=diagrams)
    buf = io.BytesIO()
    doc.save(buf)
    return Document(io.BytesIO(buf.getvalue()))


def _num_id(paragraph) -> int | None:
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    return num_pr.numId.val if num_pr is not None and num_pr.numId is not None else None


def test_word_numbered_lists_restart_and_nest():
    doc = _docx("1. first\n2. second\n\nA paragraph.\n\n" + STEPS)
    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    first_list, second_list = {_num_id(p) for p in numbered[:2]}, {_num_id(p) for p in numbered[2:]}
    assert len(first_list) == len(second_list) == 1 and first_list != second_list
    numbering = doc.part.numbering_part.element
    for num_id in first_list | second_list:
        num = numbering.num_having_numId(num_id)
        assert num.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val")) == "1"
    nested = [p.text for p in doc.paragraphs if p.style.name == "List Bullet 2"]
    assert nested == [
        "Use app-enforced restrictions: the limited experience.",
        "Block access: blocks desktop and mobile apps.",
    ]


def test_word_links_are_real_hyperlinks():
    doc = _docx("See **[the docs](https://learn.example/x)** and [local](/api/files/a.png).")
    body = doc.element.body
    links = body.findall(".//" + qn("w:hyperlink"))
    assert len(links) == 1
    rel = doc.part.rels[links[0].get(qn("r:id"))]
    assert rel.is_external and rel.target_ref == "https://learn.example/x"


def test_word_tables_keep_formatting_breaks_and_repeat_their_header():
    doc = _docx("| Method | Licence |\n|---|---|\n| **Method 1**<br>*(native)* | `P1` |\n")
    table = doc.tables[0]
    cell = table.rows[1].cells[0].paragraphs[0]
    assert [r.text for r in cell.runs if r.bold] == ["Method 1"]
    assert cell._p.findall(".//" + qn("w:br"))
    assert "<br" not in table.rows[1].cells[0].text
    assert table.rows[1].cells[1].paragraphs[0].runs[0].font.name == "Consolas"
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.rows[1]._tr.trPr.find(qn("w:cantSplit")) is not None


def test_word_code_block_is_one_shaded_paragraph_and_rules_are_borders():
    doc = _docx("```powershell\nConnect-SPOService `\n  -Url x\n```\n\n---\n")
    code = [p for p in doc.paragraphs if "Connect-SPOService" in p.text]
    assert len(code) == 1
    assert "-Url x" in code[0].text
    ppr = code[0]._p.pPr
    assert ppr.find(qn("w:shd")).get(qn("w:fill")) == "F6F8FA"
    rules = [p for p in doc.paragraphs if p._p.pPr is not None
             and p._p.pPr.find(qn("w:pBdr")) is not None and not p.text]
    assert rules and rules[-1]._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")) is not None
    assert "\u2500" not in "".join(p.text for p in doc.paragraphs)


def test_word_diagrams_are_centred_and_sized_from_their_layout():
    code = "graph TD; A-->B"
    doc = _docx(f"```mermaid\n{code}\n```", [
        {"code": code, "data": _png(1600, 1200), "width": 400, "height": 300}
    ])
    shape = doc.inline_shapes[0]
    assert shape.width.inches == pytest.approx(400 * 1.2 / 72, rel=0.01)
    assert shape.height.inches == pytest.approx(shape.width.inches * 0.75, rel=0.01)
    # A tall diagram is capped by the page height instead.
    tall = _docx(f"```mermaid\n{code}\n```", [
        {"code": code, "data": _png(1600, 3200), "width": 400, "height": 800}
    ]).inline_shapes[0]
    assert tall.height.inches == pytest.approx(8.0, rel=0.01)
    picture_para = next(p for p in doc.paragraphs if p._p.findall(".//" + qn("w:drawing")))
    assert picture_para.alignment == 1  # centre


def test_word_page_setup_matches_the_pdf():
    doc = _docx("text")
    section = doc.sections[0]
    assert section.left_margin.inches == section.right_margin.inches == 0.75
    footer = section.footer.paragraphs[0]
    assert "Title \u00b7 model" in footer._p.xml
    assert " PAGE " in footer._p.xml and " NUMPAGES " in footer._p.xml
    # The only live tab stop is the right margin, so "Page X of Y" sits flush right.
    stops = footer.paragraph_format.tab_stops
    live = [t for t in stops if t.alignment != WD_TAB_ALIGNMENT.CLEAR]
    assert [(round(t.position.inches, 2), t.alignment) for t in live] == [
        (7.0, WD_TAB_ALIGNMENT.RIGHT)
    ]
