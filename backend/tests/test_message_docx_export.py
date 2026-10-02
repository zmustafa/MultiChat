"""Offline regressions for exporting a single lane answer as a Word document."""
from __future__ import annotations

import io

import test_generated_image_ownership as ownership_fixtures
from docx import Document
from PIL import Image
from sqlalchemy import select

from app import export, models

chat = ownership_fixtures.chat
offline = ownership_fixtures.offline
owned = ownership_fixtures.owned


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "#6366f1").save(buf, format="PNG")
    return buf.getvalue()


def test_export_message_docx_contains_prompt_response_and_diagram(owned, chat):
    chat.message.content = (
        "# Findings\n\nThe **answer** is here.\n\n"
        "| a | b |\n|---|---|\n| 1 | 2 |\n\n"
        "```mermaid\ngraph TD; A-->B\n```\n"
    )
    owned.db.flush()
    diagrams = [{"code": "graph TD; A-->B", "data": _png(), "width": 40, "height": 20}]

    stored, download_name, mime = export.export_message_docx(
        owned.db, chat.session, chat.message, diagrams
    )

    assert download_name.endswith(".docx")
    assert mime == export._MIME["docx"]
    doc = Document(str(owned.root / stored))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Image ownership" in text
    assert "Render the images" in text
    assert "Response \u2014 offline" in text
    assert "The answer is here." in text
    assert "graph TD" not in text
    assert doc.tables and doc.tables[0].cell(1, 1).text == "2"
    assert doc.inline_shapes, "mermaid diagram should be embedded as an image"
    assert doc.core_properties.title == "Image ownership - offline"


def test_export_message_can_omit_prompt(owned, chat):
    from pypdf import PdfReader

    chat.message.content = "The **answer** is here."
    owned.db.flush()

    stored, _, _ = export.export_message_docx(
        owned.db, chat.session, chat.message, include_prompt=False
    )
    text = "\n".join(p.text for p in Document(str(owned.root / stored)).paragraphs)
    assert "Render the images" not in text
    assert "Prompt" not in text
    assert "The answer is here." in text

    stored, _, _ = export.export_message_pdf(
        owned.db, chat.session, chat.message, include_prompt=False
    )
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(str(owned.root / stored)).pages)
    assert "Render the images" not in pdf_text
    assert "PROMPT" not in pdf_text
    assert "The answer is here." in pdf_text

    stored, _, _ = export.export_message_pdf(owned.db, chat.session, chat.message)
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(str(owned.root / stored)).pages)
    assert "Render the images" in pdf_text


def test_message_export_route_supports_docx_and_rejects_unknown_format(client, auth, db, transcript):
    lane = sorted(transcript.lanes, key=lambda x: x.position)[0]
    message_id = db.scalars(
        select(models.LaneMessage.id).where(models.LaneMessage.lane_id == lane.id)
    ).first()
    url = f"/api/sessions/{transcript.id}/messages/{message_id}/export"

    res = client.post(f"{url}?fmt=docx", json={"diagrams": []}, headers=auth)
    assert res.status_code == 200, res.text
    assert res.json()["download_name"].endswith(".docx")
    record = db.scalars(
        select(models.GeneratedFile).where(
            models.GeneratedFile.session_id == transcript.id,
            models.GeneratedFile.kind == "docx",
        )
    ).first()
    assert record is not None and record.size_bytes > 0

    assert client.post(url, json={"diagrams": []}, headers=auth).json()[
        "download_name"
    ].endswith(".pdf")
    assert client.post(f"{url}?fmt=xlsx", json={}, headers=auth).status_code == 400
    res = client.post(
        f"{url}?fmt=docx", json={"diagrams": [], "include_prompt": False}, headers=auth
    )
    assert res.status_code == 200, res.text
